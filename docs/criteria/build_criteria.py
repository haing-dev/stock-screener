from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = "docs/criteria/Screener_Criteria.xlsx"

# shared caveats
MC = "Price x SEC share count. Unreliable for ADRs (ADR ratio unknown), multi-class shares or missing share counts -> 'check manually' (~150 ADRs last run)."
REP = "Rare: some foreign (20-F) reports filed but not yet in SEC XBRL data -> flagged."
DEBT = "Rare: debt in non-standard tags or D&A untagged -> flagged."
CAPEX = "Rare: capex not in a standard tag -> estimated or flagged."
SIC = "Uses SEC industry (SIC) code; codes can be coarse or outdated (manual override list exists)."
HUMAN = "Qualitative; not in financial data. Left to the AI review + you."
WARN = "Detected automatically; temporary vs permanent needs a human."


def hist(y, why="Need enough years to tell skill from luck.", adds="Makes return and stability numbers meaningful."):
    return ("History", f">= {y} years of data", "Data", why, adds, "Easy", "")


def recent():
    return ("Recent report", "Last annual report <= 1.6 years old", "Data",
            "Old numbers can hide a recent collapse.", "Keeps the screen current.", "Easy", REP)


def rev(x):
    return ("Size (revenue)", f"Revenue >= {x}", "Size",
            "Tiny filers have patchy data and are hard to own.", "Focus on established businesses.", "Easy", "")


def mcap(x, adds="Tradeable size."):
    return ("Size (market cap)", f"Market cap >= {x}", "Tradeability",
            "Too small to buy without moving the price.", adds, "Partly", MC)


def adv(x):
    return ("Liquidity", f"Avg daily value traded >= {x}", "Tradeability",
            "Must be able to get in and out.", "Tradeable.", "Easy", "")


def lev(x, basis="EBITDA", why="Debt can sink a good business in one bad year.", adds="Downside protection."):
    return ("Leverage", f"Net debt / {basis} <= {x}x", "Safety", why, adds, "Easy", DEBT)


def judg(what, why=HUMAN):
    return ("Human / AI review", what, "Judgment", "What numbers can't measure.",
            "Decides whether the numbers reflect a real thesis.", "No", why)


S = {}
S["Core Compounders"] = [
    ("U", "Universe", "No funds/SPACs/REITs/shells; excludes banks, insurers, real estate, commodity producers, autos, airlines, telecom, utilities, media, ad agencies", "Universe",
     "Price takers and balance-sheet businesses make ROIC/FCF misleading.", "Only businesses that can have pricing power.", "Easy", SIC),
    ("G1",) + hist(8, "A moat must show over many years, not one good stretch.", "Proves durability."),
    ("G2",) + recent(),
    ("G3",) + rev("$200M"),
    ("G4", "Consistent profits", "<= 1 operating-loss year in 10", "Quality",
     "Great businesses rarely lose money.", "Predictability.", "Easy", ""),
    ("G5a", "High returns (median)", "10-yr median ROIC >= 15%", "Quality",
     "High return on capital for years is the clearest sign of a moat.", "Moat evidence (core of the thesis).", "Easy", DEBT),
    ("G5b", "High returns (consistency)", ">= 80% of years with ROIC >= 10%", "Quality",
     "Returns must hold most years, not just on average.", "Moat is persistent, not a spike.", "Easy", ""),
    ("G6", "Growth", "10-yr revenue CAGR >= 5%", "Growth",
     "Quality without growth compounds slowly.", "Runway to reinvest.", "Easy", "Rare: tag changes cause fake revenue jumps -> flagged."),
    ("G7", "Cash conversion", "FCF / net income >= 0.8", "Quality",
     "Profits must turn into real cash.", "Earnings quality.", "Easy", CAPEX),
    ("G8", "Balance sheet", "Net debt / operating cash flow <= 3x", "Safety",
     "Too much debt can sink a good business.", "Survival / downside protection.", "Easy", DEBT),
    ("Flag", "Margin collapse", "TTM operating margin < half its 10-yr median", "Warning",
     "Might mean the moat is cracking.", "Stops buying a broken compounder.", "No", WARN),
    ("Flag", "Revenue shrinking", "3-yr revenue CAGR < 0", "Warning",
     "Cyclical dip or structural decline?", "Same.", "No", WARN),
    ("Flag", "Lost money in 2008-09", "Operating loss on falling sales in 2008-09", "Warning",
     "May be a cyclical at a good point, not a compounder.", "Avoids mistaking a cycle for a moat.", "No", WARN),
    ("V", "Valuation", "None (price-blind by design); P/E, EV/EBITDA, FCF yield and peer medians shown only", "Valuation",
     "Quality first; price decides timing, not inclusion.", "Keeps great businesses on the list even when dear.", "n/a", "No 'cheap enough' test - see Valuation sheet."),
    ("J",) + judg("Moat durability, switching costs, management and capital allocation, AI disruption, why mispriced"),
]
S["Emerging Compounders"] = [
    ("U", "Universe", "Core universe, also minus commodity materials (chemicals, paper, building materials, non-metal mining)", "Universe",
     "Commodity materials don't compound like niche leaders.", "Only potential niche winners.", "Easy", SIC),
    ("E1",) + hist(4, "Need a few years to see growth and margin trends.", "Makes trend rules meaningful."),
    ("E2",) + recent(),
    ("E3", "Already profitable", "TTM and last-FY EBITDA > 0", "Quality",
     "88% of multibaggers were already EBITDA-positive.", "Cuts cash-burning stories.", "Easy", "Rare: D&A untagged -> EBIT used."),
    ("E4", "Fast growth", "3-yr revenue CAGR >= 15%", "Growth",
     "Typical multibagger grew ~20%/yr.", "The growth engine.", "Easy", ""),
    ("E5", "Growth not fading", "Latest year-to-date growth >= 8% (else last FY)", "Growth",
     "Growth must not be slowing now.", "Catches slowdowns early.", "Partly", "Foreign filers file no quarterly reports -> uses last fiscal year (stale)."),
    ("E6", "Gross margin", "Gross margin >= 35%", "Quality",
     "Room for profit to grow faster than sales.", "Operating-leverage potential.", "Easy", "Rare: no gross-profit line (services) -> no data."),
    ("E7", "Operating leverage", "EBITDA margin up over 3 yrs, or SG&A/revenue down", "Quality",
     "Costs should grow slower than sales.", "Proof the model scales.", "Easy", ""),
    ("E8", "Little dilution", "Share count growth <= 3%/yr (3 yrs)", "Quality",
     "Dilution eats per-share gains.", "Per-share compounding.", "Partly", "No share history for recent IPOs/spin-offs and some foreign filers (~20 flagged)."),
    ("E9",) + lev(2.5, why="Small companies can't carry much debt."),
    ("E10",) + mcap("$150M", "Micro caps were the study's worst bucket."),
    ("E11", "Not already large", "Market cap <= $10B", "Size",
     "Must be small enough to still multiply.", "Runway for a multibagger.", "Partly", MC),
    ("E12", "Valuation", "EV / EBITDA <= 20x", "Valuation",
     "Winners started ~10x, mostly under 20x.", "Return comes from growth, not multiple expansion.", "Partly", MC),
    ("E13",) + adv("$200k"),
    ("J",) + judg("Why overlooked, niche dominance, customer concentration (~70% = kill), insider stake >= 5%, acquisition-led growth"),
]
S["Special Situations"] = [
    ("U", "Universe (event)", "Spin-off (Form 10), bankruptcy (8-K item 1.03) or IPO (S-1/F-1) in last 36 months; no funds/banks/insurers", "Trigger",
     "Forced sellers and no analyst coverage cause mispricing.", "The source of the edge.", "Partly", "Spin-offs without a Form 10 (e.g. foreign) can be missed; 8-K 1.03 dates the filing, not the emergence."),
    ("S1", "Recent report", "Last annual report <= 1.6 years old", "Data",
     "Need current numbers.", "Keeps the screen current.", "Partly", "New listings often have no annual report yet (only Form 10/prospectus) -> no data."),
    ("S2", "Profitable business", "TTM operating income > 0", "Quality",
     "Events work best in good businesses.", "Avoids broken companies.", "Partly", "New listings often lack tagged history -> no data."),
    ("S3",) + lev(4, why="Spin-offs often get loaded with debt.", adds="Equity isn't an option bet."),
    ("S4", "Cheap", "EV / EBIT <= 15x", "Valuation",
     "Pay a low pre-tax multiple (10x EBIT ~ a bond yield).", "Margin of safety on the mispricing.", "Partly", MC + " New listings may lack a share count."),
    ("S5",) + mcap("$100M"),
    ("S6",) + adv("$200k"),
    ("J",) + judg("Why the parent spins it off, who must sell, management incentives, debt loaded on the spinco, hidden assets", "Needs reading the Form 10 / reorganization plan."),
]
S["Fortress"] = [
    ("U", "Universe", "Same as Core", "Universe", "Same as Core.", "Same as Core.", "Easy", SIC),
    ("F1",) + hist(9, "Record must include 2020.", "Lets the recession tests run."),
    ("F2",) + recent(),
    ("F3",) + rev("$500M"),
    ("F4", "Held up in 2020", "2020 revenue >= -5% and operating income >= -15%", "Safety",
     "Real resilience shows in a recession.", "Recession-proof earnings (the thesis).", "Easy", ""),
    ("F5", "Held up in 2009", "2009 revenue >= -10%, operating income >= -25%; recession loss fails", "Safety",
     "Same test, deeper recession.", "Second proof of resilience.", "Partly", "SEC XBRL data mostly starts 2009-11 -> 2008-09 missing for many (~35 flagged); check old annual reports."),
    ("F6", "Never lost money", "0 operating-loss years", "Quality",
     "Fortress businesses don't lose money.", "Predictability.", "Easy", ""),
    ("F7", "Stable margins", "Operating margin std. dev. <= 5 pts (10 yrs)", "Quality",
     "Stable margins = predictable earnings.", "Low earnings volatility.", "Easy", ""),
    ("F8", "Decent returns", "10-yr median ROIC >= 10%", "Quality",
     "Decent returns; lower bar lets rail/waste qualify.", "Quality floor.", "Easy", DEBT),
    ("F9",) + lev(2, why="Survive first."),
    ("F10", "Interest coverage", "EBIT / interest >= 6x", "Safety",
     "Debt must stay affordable when it rolls at higher rates.", "Refinancing safety.", "Easy", "If interest isn't tagged, a worst-case estimate is used."),
    ("F11", "Refinancing", "Debt due in 2 yrs <= cash + 2 yrs of FCF", "Safety",
     "Maturities beyond cash can trigger a crisis.", "No liquidity squeeze.", "Partly", "Maturity schedule often not tagged (e.g. LMT, KMB) -> check the debt note."),
    ("F12", "Cash every year", "<= 1 negative-FCF year", "Quality",
     "Cash every year, recessions included.", "Resilient cash flow.", "Easy", CAPEX),
    ("F13", "Inflation test", "Gross margin change 2021-23 >= -2 pts", "Quality",
     "Kept margins through inflation = pricing power.", "Pricing-power evidence.", "Easy", ""),
    ("Flag", "Margin collapse / revenue shrinking", "As in Core", "Warning", "Resilience may be breaking.", "Early warning.", "No", WARN),
    ("V", "Valuation", "None (price-blind by design)", "Valuation",
     "Resilience first; price decides timing.", "-", "n/a", "No 'cheap enough' test - see Valuation sheet."),
    ("J",) + judg("Is resilience structural (essential, recurring)? Contract durability, regulation, what would break it"),
]
S["Healthcare - Commercial"] = [
    ("U", "Universe", "Healthcare SIC codes: pharma, devices, tools, distribution, insurers, services", "Universe",
     "Healthcare has its own economics (patents, payers).", "Like-for-like comparison.", "Easy", SIC),
    ("H1",) + hist(5),
    ("H2",) + recent(),
    ("H3",) + rev("$200M"),
    ("H4", "Consistent profits", "<= 1 loss year in last 5", "Quality",
     "Commercial-stage means reliably profitable.", "Predictability.", "Easy", ""),
    ("H5", "High returns", "10-yr median ROIC >= 12%", "Quality",
     "Healthcare winners earn persistently high returns.", "Moat evidence (patents, installed base).", "Easy", DEBT),
    ("H6", "Gross margin (products)", "Gross margin >= 50%", "Quality",
     "Real products earn high margins; low = distributor/commodity.", "Pricing power.", "Easy", "Rare: no gross-profit line (e.g. REGN) -> no data."),
    ("H7", "Cash conversion", "FCF / net income >= 0.7", "Quality",
     "Profits must turn into cash.", "Earnings quality.", "Easy", CAPEX),
    ("H8",) + lev(3, why="Avoid levered roll-ups."),
    ("H9", "Growing", "Revenue CAGR >= 3%", "Growth",
     "Patent cliffs show up as shrinking sales.", "Franchise still growing.", "Easy", ""),
    ("Flag", "Margin collapse / revenue shrinking", "As in Core", "Warning", "Possible patent cliff or price cut.", "Early warning.", "No", WARN),
    ("V", "Valuation", "None (price-blind by design)", "Valuation", "Quality first.", "-", "n/a", "No 'cheap enough' test - see Valuation sheet."),
    ("J",) + judg("Patent expiries next 5 yrs (% revenue), Medicare/Medicaid share, price-led growth, pipeline depth", "Not in financial data; AI review + 10-K reading."),
]
S["Healthcare - Clinical"] = [
    ("U", "Universe", "Pharma/biotech/device SIC codes, not commercial-stage (revenue < $200M or operating loss)", "Universe",
     "Pre-revenue firms can't be judged on earnings.", "Separate 'cash floor + option' logic.", "Easy", SIC),
    ("B1", "Recent balance sheet", "Balance sheet <= 200 days old", "Data",
     "Cash burns fast; the figure must be fresh.", "Accurate floor.", "Easy", ""),
    ("B2", "Cash runway", "Cash / annual burn >= 2 yrs", "Safety",
     "Avoid forced share sales at the bottom.", "Time to reach the next catalyst.", "Easy", "Burn = last 12 months operating cash flow."),
    ("B3", "Little debt", "Debt / cash <= 0.5", "Safety",
     "Debt and converts eat the cash floor.", "Floor belongs to shareholders.", "Easy", ""),
    ("B4",) + mcap("$50M"),
    ("B5",) + adv("$200k"),
    ("B6", "Floor value", "EV / cash <= 0.25", "Valuation",
     "Pay little more than the cash; the science is a cheap option.", "Hard downside floor + free upside.", "Partly", MC),
    ("J",) + judg("Lead asset and phase, next catalyst and date, science quality, crowded mechanism, management record", "Numbers can't judge science; AI review reads the pipeline."),
]
S["Cyclicals"] = [
    ("U", "Universe", "Capacity-cycle industries: metals, energy, chemicals, building materials, homebuilders, machinery, semis, autos, transport", "Universe",
     "These earn in cycles, so need cycle-aware rules.", "Right yardstick for cyclicals.", "Easy", SIC),
    ("C1",) + hist(8, "Need a full cycle to normalize earnings.", "Enables cycle-average figures."),
    ("C2",) + recent(),
    ("C3",) + rev("$200M"),
    ("C4", "Mid-cycle returns", "10-yr median ROIC >= 10%", "Quality",
     "A good business even in a cyclical industry.", "Quality filter.", "Easy", DEBT),
    ("C5", "Survives downturns", ">= 7 profitable (EBIT) years", "Quality",
     "Must stay profitable through most of the cycle.", "Survival.", "Easy", ""),
    ("C6",) + lev(2.5, "avg EBITDA (cycle)", "A trough mustn't break it.", "Survives the next downturn."),
    ("C7", "Not at peak margins", "Latest margin in lower half of own 10-yr range", "Trigger",
     "Buy when margins are depressed, not at the peak.", "Cycle timing (core of the thesis).", "Easy", "Own history only; industry capacity cycle isn't measured."),
    ("C8", "Cheap on normalized earnings", "EV / 10-yr avg EBIT <= 12x", "Valuation",
     "Cheap on cycle-average earnings.", "Margin of safety if the cycle turns slowly.", "Partly", MC),
    ("J",) + judg("Industry capacity (closures vs expansions), cyclical vs structural margin drop", "Industry capacity data isn't in SEC filings."),
]
S["Owner-Operators"] = [
    ("U", "Universe", "Same as Core", "Universe", "Same as Core.", "Same as Core.", "Easy", SIC),
    ("O1",) + hist(5),
    ("O2",) + recent(),
    ("O3",) + rev("$50M"),
    ("O4", "Insider ownership", "Insider share >= 20%", "Trigger",
     "Owners who run the company act like owners.", "Alignment + under-owned by index funds (cheaper).", "Partly",
     "From cover-page public float. Foreign filers report none (~43); no share count near float date (~29); share-class mismatches."),
    ("O5", "Consistent profits", "<= 2 operating-loss years", "Quality",
     "A real business, not a lottery ticket.", "Quality floor.", "Easy", ""),
    ("O6", "High returns", "10-yr median ROIC >= 12%", "Quality",
     "Owners compound only if the business earns well.", "Compounding engine.", "Easy", DEBT),
    ("O7", "Cash conversion", "FCF / net income >= 0.6", "Quality",
     "Profits must turn into cash.", "Earnings quality.", "Easy", CAPEX),
    ("O8",) + lev(2.5, why="Owners who borrow heavily can lose the company."),
    ("O9",) + mcap("$100M"),
    ("O10",) + adv("$100k"),
    ("O11", "Valuation", "EV / EBIT <= 20x", "Valuation",
     "Owner-operators are often cheap; don't overpay.", "Entry discount the thesis relies on.", "Partly", MC),
    ("Flag", "Margin collapse / revenue shrinking", "As in Core", "Warning", "Business may be weakening.", "Early warning.", "No", WARN),
    ("J",) + judg("Founder/family vs parent company, pay structure, related-party deals, voting classes, succession", "Needs the proxy statement (DEF 14A); not parsed."),
]
S["Hard Assets & Royalties"] = [
    ("U", "Universe", "Land, mineral/royalty, oil & gas, timber, pipeline and water owners (no REITs)", "Universe",
     "Asset owners, not operators.", "Where royalty-like assets sit.", "Easy", SIC + " R4 separates owners from producers."),
    ("R1",) + hist(5, "Enough years to see a cycle."),
    ("R2",) + recent(),
    ("R3",) + rev("$25M"),
    ("R4", "Little reinvestment", "Capex / operating cash flow (3 yrs) <= 20%", "Trigger",
     "True asset owners barely reinvest.", "Separates royalty owners from producers.", "Easy", CAPEX),
    ("R5", "Royalty-like margins", "10-yr median operating margin >= 30%", "Quality",
     "Owning, not working, the asset gives very high margins.", "Confirms owner economics.", "Easy", ""),
    ("R6", "Consistent profits", "<= 2 operating-loss years", "Quality",
     "Profitable through commodity cycles.", "Resilience.", "Easy", ""),
    ("R7",) + lev(2, why="A commodity slump mustn't break it."),
    ("R8",) + mcap("$100M"),
    ("R9",) + adv("$100k"),
    ("R10", "Valuation", "EV / EBIT <= 25x", "Valuation",
     "Even a perfect royalty can be bought too dear.", "Avoids overpaying.", "Partly", MC),
    ("J",) + judg("Dormant assets not on the books, single-commodity dependence, trust end dates, operator quality"),
]
S["Asset-Floor Value"] = [
    ("U", "Universe", "All except funds/SPACs/REITs, banks, insurers, real estate", "Universe",
     "Cash/book floors don't apply to balance-sheet businesses.", "Broad hunt for deep value.", "Easy", "BDCs/lenders slip in -> often 'no data'."),
    ("K1",) + hist(3, "Enough years to see a record."),
    ("K2",) + recent(),
    ("K3", "Recent balance sheet", "Balance sheet <= 200 days old", "Data",
     "The floor is measured on it.", "Accurate floor.", "Easy", ""),
    ("K4", "Below a hard floor", "Market cap <= 1.0x net cash, or price <= 1.0x tangible book", "Valuation",
     "Buy below cash and get the business free.", "Hard downside floor (the thesis).", "Partly",
     MC + " Tangible book ignores off-balance-sheet liabilities."),
    ("K5", "Not burning the floor", "Operating cash flow >= 0, or >= 4 yrs of cash", "Safety",
     "A floor being burned isn't a floor.", "Floor stays intact.", "Easy", ""),
    ("K6", "Usually profitable", "Profitable in >= 60% of years", "Quality",
     "A cheap business that never earns is a value trap.", "Avoids value traps.", "Easy", ""),
    ("K7",) + lev(3, why="Debt ahead of shareholders erodes the floor."),
    ("K8",) + mcap("$50M"),
    ("K9",) + adv("$100k"),
    ("J",) + judg("Will management waste the cash? Cash trapped abroad / VIE? Real inventories? Off-balance-sheet items? Catalyst?"),
]
S["Trough Cyclical Cash"] = [
    ("U", "Universe", "Same as Cyclicals", "Universe", "Same as Cyclicals.", "Same as Cyclicals.", "Easy", SIC),
    ("T1",) + hist(7, "Need a full cycle to judge."),
    ("T2",) + recent(),
    ("T3",) + rev("$200M"),
    ("T4", "Margins below normal", "Latest margin in lower half of own 10-yr range", "Trigger",
     "Buy at the trough, not mid-cycle.", "Cycle timing.", "Easy", ""),
    ("T5", "Cash through the cycle", "<= 3 negative-FCF years and 5-yr avg FCF > 0", "Quality",
     "Pays you to wait, even at the trough.", "Cash machine through the cycle.", "Partly", "Energy producers often tag capex non-standard -> too few FCF years (~8 flagged)."),
    ("T6", "Not adding capacity", "Capex / D&A (3 yrs) <= 1.3", "Quality",
     "Maintaining, not expanding = supply discipline.", "Supply side of the thesis.", "Partly", "Capex or D&A untagged for some producers (~8 flagged)."),
    ("T7",) + lev(2.5, "avg EBITDA (cycle)", "A trough mustn't break it."),
    ("T8", "Cash yield on the cycle", "5-yr avg FCF / market cap >= 10%", "Valuation",
     "Cycle-average cash yield must be high.", "Cheap on through-cycle cash.", "Partly", MC + " Needs 5 yrs of FCF."),
    ("T9",) + mcap("$200M"),
    ("T10",) + adv("$200k"),
    ("Rank", "Industry returning cash", "Industry buybacks + dividends / capex (ranks only)", "Ranking",
     "Industries returning cash instead of building rank higher.", "Supply discipline at industry level.", "Partly", "Industry = SIC group; coarse."),
    ("J",) + judg("Trough vs mid-cycle, supply leaving the industry, bear case vs revenue mix, freight/regulation protection"),
]
S["Hidden Earnings"] = [
    ("U", "Universe", "Same as Asset-Floor", "Universe", "Broad hunt.", "Broad hunt.", "Easy", ""),
    ("X1",) + hist(5, "Enough years to know the business."),
    ("X2",) + recent(),
    ("X3", "Something the market is ignoring", "Discontinued ops >= 1% of revenue (3 yrs), or price >= 40% below 12-mo high", "Trigger",
     "Market fixates on a problem; a sale/spin-off lifts the rest.", "The mispricing trigger.", "Partly",
     "Discontinued-ops tags used inconsistently; 12-mo high needs stored price history + split list (else no data)."),
    ("X4", "Still earning", "TTM operating income > 0", "Quality",
     "The remaining business must be profitable now.", "Not a broken company.", "Easy", ""),
    ("X5", "Sales holding up", "3-yr revenue CAGR >= -5%", "Growth",
     "Temporary problem, not steady decline.", "Separates dip from decline.", "Easy", ""),
    ("X6", "Still generating cash", "TTM FCF > 0", "Quality",
     "Cash keeps coming in.", "Earnings quality.", "Easy", CAPEX),
    ("X7", "Decent returns", "10-yr median ROIC >= 8%", "Quality",
     "A good business behind the mess.", "Quality floor.", "Easy", DEBT),
    ("X8",) + lev(3, why="Debt turns a messy situation into an option bet."),
    ("X9", "Cheap", "EV / EBIT <= 15x", "Valuation",
     "Messy situations work best bought cheaply.", "Margin of safety.", "Partly", MC),
    ("X10",) + mcap("$100M"),
    ("X11",) + adv("$200k"),
    ("J",) + judg("Which segment is feared, revenue elsewhere, parent ex-unit, why price fell, who is forced to sell", "Needs segment notes; not parsed."),
]
S["Red Flags (avoid)"] = [
    ("U", "Universe", "Same as Asset-Floor, minus clinical-stage biotech", "Universe",
     "Biotechs burn cash by design.", "Avoids false alarms.", "Easy", ""),
    ("W1",) + hist(3, "Enough years to see a trend."),
    ("W2",) + recent(),
    ("W3", "Priced for perfection", "EV / sales >= 6x", "Valuation",
     "Expensive names de-rate hard on any miss.", "First half of the avoid case.", "Partly", MC),
    ("W4", "Something is breaking", ">= 2 of: growth < half its 3-yr rate; gross margin down 3 yrs; op margin down >= 5 pts; TTM loss; dilution >= 5%/yr; shrinking revenue", "Trigger",
     "Expensive + deteriorating = the classic short.", "Second half of the avoid case.", "Partly", "Needs 3 yrs of tagged history (~28 can't be judged)."),
    ("W5",) + mcap("$500M", "Big enough to matter."),
    ("W6",) + adv("$1M"),
    ("J",) + judg("One-off vs real deterioration, short interest/squeeze risk, AI disruption, insider selling", "Short interest and insider-trade data not connected."),
]

VAL = [
    ("Core Compounders", "None (price-blind)", "-"),
    ("Emerging Compounders", "EV/EBITDA <= 20x", "Absolute multiple"),
    ("Special Situations", "EV/EBIT <= 15x", "Absolute multiple"),
    ("Fortress", "None (price-blind)", "-"),
    ("Healthcare - Commercial", "None (price-blind)", "-"),
    ("Healthcare - Clinical", "EV/cash <= 0.25", "Asset floor"),
    ("Cyclicals", "EV / 10-yr avg EBIT <= 12x", "Absolute, on normalized earnings"),
    ("Owner-Operators", "EV/EBIT <= 20x", "Absolute multiple"),
    ("Hard Assets & Royalties", "EV/EBIT <= 25x", "Absolute multiple"),
    ("Asset-Floor Value", "Mcap <= net cash, or P/TBV <= 1.0x", "Asset floor"),
    ("Trough Cyclical Cash", "5-yr avg FCF yield >= 10%", "Absolute, on cycle-average cash"),
    ("Hidden Earnings", "EV/EBIT <= 15x", "Absolute multiple"),
    ("Red Flags (avoid)", "EV/sales >= 6x (expensive = flag)", "Absolute multiple"),
]

import json as _j, glob as _g
RAT = {}
for _f in _g.glob("docs/criteria/out_*.json"):
    RAT.update(_j.load(open(_f)))
for _n, _rows in S.items():
    S[_n] = [r[:4] + (RAT[f"{_n}|{r[0]}|{r[1]}"],) + r[5:] for r in _rows]
ARIAL = "Arial"
hdr_font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", start_color="1F3864")
body = Font(name=ARIAL, size=9)
bold = Font(name=ARIAL, size=9, bold=True)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(wrap_text=True, vertical="top", horizontal="center")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
FILLS = {"Easy": "E2EFDA", "Partly": "FFF2CC", "No": "F8CBAD", "n/a": "EDEDED"}


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, center, border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28


wb = Workbook()
ws = wb.active
ws.title = "Criteria"
cols = ["Screen", "ID", "Criterion", "Threshold", "Type", "Why it's there (investor rationale)", "What it adds to the thesis",
        "Verifiable now?", "Why not easy / caveat"]
header(ws, cols, [18, 6, 20, 32, 11, 90, 30, 10, 44])
r = 2
for screen, rows in S.items():
    for row in rows:
        vals = [screen] + list(row)
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = bold if c in (1, 3) else body
            cell.alignment = center if c in (2, 5, 8) else wrap
            cell.border = border
        r += 1
last = r - 1
ws.auto_filter.ref = f"A1:I{last}"
for k, color in FILLS.items():
    ws.conditional_formatting.add(f"H2:H{last}", FormulaRule(formula=[f'$H2="{k}"'], fill=PatternFill("solid", start_color=color)))
ws.conditional_formatting.add(f"A2:I{last}", FormulaRule(formula=['$E2="Valuation"'], font=Font(name=ARIAL, color="1F3864", bold=True)))

# Valuation sheet
wv = wb.create_sheet("Valuation")
vcols = ["Screen", "'Cheap enough' rule (pass/fail)", "Kind", "Vs peers", "Vs own history", "Implied return / DCF"]
header(wv, vcols, [24, 34, 28, 30, 26, 20])
for i, (s, rule, kind) in enumerate(VAL, 2):
    vals = [s, rule, kind, "Shown on site (industry median, green/red); not a rule", "No", "No"]
    for c, v in enumerate(vals, 1):
        cell = wv.cell(row=i, column=c, value=v)
        cell.font = bold if c == 1 else body
        cell.alignment, cell.border = wrap, border
n = len(VAL) + 3
notes = [
    "Answer: 9 of 13 screens have an absolute 'cheap enough' rule. None compares price to peers or to the company's own past multiples as a pass/fail rule.",
    "Core, Fortress and HC Commercial are price-blind by design: they show P/E, EV/EBITDA, FCF yield and peer medians, but never exclude on price.",
    "Possible adds: (1) peer-relative rule, e.g. EV/EBIT <= industry median (data already computed); (2) FCF-yield hurdle for Core/Fortress/HC; (3) own 5-10 yr multiple percentile (needs multi-year price history, not stored now).",
]
for j, t in enumerate(notes):
    cell = wv.cell(row=n + j, column=1, value=t)
    cell.font = bold if j == 0 else body
    wv.merge_cells(start_row=n + j, start_column=1, end_row=n + j, end_column=6)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    wv.row_dimensions[n + j].height = 30

# Summary sheet (formulas)
sm = wb.create_sheet("Summary")
scols = ["Screen", "Easy", "Partly", "No", "n/a", "Total", "% Easy"]
header(sm, scols, [26, 9, 9, 9, 9, 9, 10])
names = list(S.keys())
for i, s in enumerate(names, 2):
    sm.cell(row=i, column=1, value=s).font = bold
    for c, k in zip(range(2, 6), ["Easy", "Partly", "No", "n/a"]):
        sm.cell(row=i, column=c, value=f'=COUNTIFS(Criteria!$A$2:$A${last},$A{i},Criteria!$H$2:$H${last},"{k}")')
    sm.cell(row=i, column=6, value=f"=SUM(B{i}:E{i})")
    sm.cell(row=i, column=7, value=f"=IF(F{i}=0,0,B{i}/F{i})").number_format = "0%"
tr = len(names) + 2
sm.cell(row=tr, column=1, value="All screens").font = bold
for c in range(2, 7):
    L = get_column_letter(c)
    sm.cell(row=tr, column=c, value=f"=SUM({L}2:{L}{tr-1})").font = bold
sm.cell(row=tr, column=7, value=f"=IF(F{tr}=0,0,B{tr}/F{tr})").number_format = "0%"
for row in sm.iter_rows(min_row=2, max_row=tr):
    for cell in row:
        if cell.column > 1:
            cell.font = Font(name=ARIAL, size=9, bold=cell.row == tr)
            cell.alignment = Alignment(horizontal="center")
        cell.border = border
leg = tr + 2
legend = [
    ("Legend", None),
    ("Easy", "Checked automatically from SEC XBRL / price data and reliable for nearly all companies; rare exceptions are auto-flagged."),
    ("Partly", "Automatic for most, but a known group (ADRs, multi-class shares, foreign filers, untagged items) lands in 'check manually'."),
    ("No", "Not checked by numbers: warning flags or judgment items for the AI review + you."),
    ("n/a", "No rule (price-blind screens)."),
    ("Rationale", "Format: investor states that X means Y, because Z; example. Checked against all 146 transcripts. 'No investor source' = practical data rule."),
    ("Note", "Each row is in config.json; thresholds can be changed there. Counts like '~150' are from the 7 Oct 2026 check-manually report."),
]
for j, (k, t) in enumerate(legend):
    a = sm.cell(row=leg + j, column=1, value=k)
    a.font = bold
    if k in FILLS:
        a.fill = PatternFill("solid", start_color=FILLS[k])
    if t:
        b = sm.cell(row=leg + j, column=2, value=t)
        b.font = body
        sm.merge_cells(start_row=leg + j, start_column=2, end_row=leg + j, end_column=7)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        sm.row_dimensions[leg + j].height = 26

wb.save(OUT)
print(OUT, "rows", last - 1)
